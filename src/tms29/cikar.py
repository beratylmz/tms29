"""Excel çalışma kitabından kayıpsız veri çıkarımı.

Tasarım kararları
-----------------
1. **Uzun (tidy) biçim, geniş değil.** Bir satır = (şirket, tablo, ham
   etiket, yıl, fiyat seviyesi, blok, değer). Şirketler farklı satır
   düzenleri kullandığı için geniş biçim ya çok seyrek bir matris ya da
   erken bir yorum gerektirirdi. Uzun biçimde düzen farkı sorun değil.

2. **Bu katman yorum yapmaz.** Etiketler ham hâliyle taşınıyor;
   kanonik isme çevirme işi `kalemler.py`'de, ayrı ve test edilebilir
   bir katmanda. Çıkarım ile yorumu aynı fonksiyona koymak, veri
   kaybının fark edilmediği yerdir.

3. **Kolon bloğu değerden anlaşılıyor, başlıktan değil.** Kaynak
   kitapta her tablo iki blok hâlinde: mutlak tutarlar ve dikey yüzde
   (common-size). İki bloğun başlık satırı birebir aynı. Başlığa göre
   tekilleştirseydik ikinci bloğun varlığını hiç fark etmezdik --
   ve blok sırası bir sayfada ters olsaydı yüzdeleri TL sanardık.
   Ayrım `DIKEY_YUZDE_TAVANI` ile yapılıyor (bkz. kavramlar.py).

4. **Fiyat seviyesi asla yıl ile birleştirilmiyor.** Projenin ölçtüğü
   şey tam olarak aynı yılın iki fiyat seviyesi arasındaki fark.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from tms29.kavramlar import DIKEY_YUZDE_TAVANI, fiyat_seviyesi_ayristir

#: Sayfa adı sonekleri -> tablo adı.
TABLOLAR = {"_BS": "bilanco", "_PLS": "gelir", "_CFS": "nakit_akis"}


def _baslik_satiri(ws) -> int:
    """Yıl içeren ilk satırı başlık kabul eder.

    Sayfalar bazen bir başlık bandıyla başlıyor, bazen doğrudan
    kolonlarla. Sabit satır numarası varsaymak bu yüzden kırılgan.
    """
    for r in range(1, 7):
        sayac = sum(
            1
            for c in range(2, ws.max_column + 1)
            if "20" in str(ws.cell(row=r, column=c).value or "")
        )
        if sayac >= 2:
            return r
    raise ValueError(f"{ws.title}: başlık satırı bulunamadı")


def _blok_tipi(ws, kolon: int, ilk_veri_satiri: int) -> str:
    """Kolonun mutlak tutar mı yoksa dikey yüzde mi olduğunu söyler."""
    en_buyuk = 0.0
    for r in range(ilk_veri_satiri, ws.max_row + 1):
        v = ws.cell(row=r, column=kolon).value
        if isinstance(v, (int, float)):
            en_buyuk = max(en_buyuk, abs(float(v)))
    if en_buyuk == 0.0:
        return "bos"
    return "dikey_yuzde" if en_buyuk <= DIKEY_YUZDE_TAVANI else "mutlak"


def calisma_kitabini_cikar(yol: str | Path) -> pd.DataFrame:
    """Bütün sayfaları tek bir uzun DataFrame'e çıkarır.

    Returns
    -------
    DataFrame
        Kolonlar: sirket, tablo, ham_etiket, yil, seviye, blok, deger
    """
    import openpyxl

    wb = openpyxl.load_workbook(Path(yol), data_only=True)
    satirlar: list[dict] = []

    for sayfa in wb.sheetnames:
        sonek = next((s for s in TABLOLAR if sayfa.endswith(s)), None)
        if sonek is None:
            continue  # Ratios / Analysis / karşılaştırma sayfaları: türetilmiş
        ws = wb[sayfa]
        sirket = sayfa[: -len(sonek)]
        tablo = TABLOLAR[sonek]
        bs = _baslik_satiri(ws)

        kolonlar = []
        for c in range(2, ws.max_column + 1):
            h = str(ws.cell(row=bs, column=c).value or "").replace("\n", " ").strip()
            if not h:
                continue
            try:
                yil, seviye = fiyat_seviyesi_ayristir(h)
            except ValueError:
                continue  # yıl/seviye taşımayan yardımcı kolon
            blok = _blok_tipi(ws, c, bs + 1)
            if blok == "bos":
                continue
            kolonlar.append((c, yil, seviye, blok))

        for r in range(bs + 1, ws.max_row + 1):
            etiket = str(ws.cell(row=r, column=1).value or "").strip()
            if not etiket:
                continue
            for c, yil, seviye, blok in kolonlar:
                v = ws.cell(row=r, column=c).value
                if not isinstance(v, (int, float)):
                    continue
                satirlar.append(
                    dict(
                        sirket=sirket,
                        tablo=tablo,
                        ham_etiket=etiket,
                        yil=yil,
                        seviye=seviye,
                        blok=blok,
                        deger=float(v),
                    )
                )

    df = pd.DataFrame(satirlar)
    # Aynı (şirket, tablo, etiket, yıl, seviye, blok) birden fazla kolonda
    # görünebiliyor -- kaynakta bazı kolonlar gerçekten kopyalanmış.
    # Kopyaları düşürmeden önce çelişki var mı diye bakılıyor.
    anahtar = ["sirket", "tablo", "ham_etiket", "yil", "seviye", "blok"]
    celiski = df.groupby(anahtar)["deger"].nunique()
    celiskili = celiski[celiski > 1]
    if len(celiskili):
        raise ValueError(
            f"Aynı anahtarda çelişen değerler var ({len(celiskili)} adet). "
            f"İlk örnek: {celiskili.index[0]}"
        )
    return df.drop_duplicates(subset=anahtar).reset_index(drop=True)
