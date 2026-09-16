"""Çıkarım testleri — paketin en riskli kodu burası.

Kaynak çalışma kitabındaki tuzak: her tablo iki blok hâlinde duruyor,
önce bin TL cinsinden mutlak tutarlar, sonra aynı kalemlerin toplama
oranı (dikey yüzde). **İki bloğun başlık satırı birebir aynı.**

Bu fark edilmeseydi ne olurdu: başlığa göre tekilleştirme yapılır,
ikinci bloğun varlığı hiç görülmez; bir sayfada blok sırası ters olsa
yüzdeler TL sanılır ve bütün analiz sessizce saçmalardı. Aşağıdaki ilk
test bunu kilitliyor.

Testler kaynak Excel'i gerçekten okuyor (openpyxl dev bağımlılığı).
Dondurulmuş CSV'nin kaynaktan yeniden üretilebildiğini de doğruluyor --
yoksa veri dosyası zamanla koddan kopar ve kimse fark etmez.
"""

import pandas as pd
import pytest

from tms29.cikar import calisma_kitabini_cikar
from tms29.veri import KAYNAK_YOL, ham_yukle

pytest.importorskip("openpyxl")


@pytest.fixture(scope="module")
def cikarim():
    if not KAYNAK_YOL.exists():  # pragma: no cover
        pytest.skip("kaynak çalışma kitabı depoda yok")
    return calisma_kitabini_cikar(KAYNAK_YOL)


# --- BLOK AYRIMI: bu dosyanın varlık sebebi --------------------------------


def test_iki_blok_da_yakalandi(cikarim):
    assert set(cikarim.blok) == {"mutlak", "dikey_yuzde"}


def test_hasilat_iki_blokta_farkli_deger_tasiyor(cikarim):
    """ASTOR 2025 hasılatı: mutlak blokta milyarlar, yüzde blokta 1,0.

    Aynı başlık, aynı kalem, iki ayrı anlam. Blok etiketi olmasaydı
    ayırt edilemezlerdi.
    """
    d = cikarim[
        (cikarim.sirket == "ASTOR")
        & (cikarim.tablo == "gelir")
        & (cikarim.yil == 2025)
        & (cikarim.ham_etiket.str.startswith("Hasılat"))
    ]
    mutlak = d[d.blok == "mutlak"].deger.iloc[0]
    yuzde = d[d.blok == "dikey_yuzde"].deger.iloc[0]
    assert mutlak > 1e6  # bin TL cinsinden milyarlarca
    assert yuzde == pytest.approx(1.0, abs=0.01)  # hasılat / hasılat


def test_dikey_yuzde_blogu_tavani_asmiyor(cikarim):
    """Sınıflandırmanın tanımı: yüzde bloğunda büyük tutar bulunmaz."""
    from tms29.kavramlar import DIKEY_YUZDE_TAVANI

    yuzde = cikarim[cikarim.blok == "dikey_yuzde"]
    assert yuzde.deger.abs().max() <= DIKEY_YUZDE_TAVANI


def test_mutlak_blok_gercek_tutarlar_iceriyor(cikarim):
    mutlak = cikarim[cikarim.blok == "mutlak"]
    assert mutlak.deger.abs().max() > 1e8


# --- KAYIPSIZLIK VE TEKRARLANABİLİRLİK ------------------------------------


def test_dondurulmus_veri_kaynaktan_yeniden_uretilebiliyor(cikarim):
    """Depodaki CSV, kaynak Excel'den birebir çıkıyor mu?

    Çıkmıyorsa veri dosyası koddan kopmuş demektir.
    """
    anahtar = ["sirket", "tablo", "ham_etiket", "yil", "seviye", "blok"]
    a = cikarim.sort_values(anahtar).reset_index(drop=True)
    b = ham_yukle().sort_values(anahtar).reset_index(drop=True)
    assert len(a) == len(b)
    pd.testing.assert_frame_equal(a[anahtar], b[anahtar])
    assert (a.deger - b.deger).abs().max() < 1e-6


def test_yedi_sirket_ve_uc_tablo(cikarim):
    assert len(set(cikarim.sirket)) == 7
    assert set(cikarim.tablo) == {"bilanco", "gelir", "nakit_akis"}


def test_fiyat_seviyeleri_ayri_kaldi(cikarim):
    """2022 hem nominal hem düzeltilmiş olarak duruyor -- ölçümün dayanağı."""
    bs = cikarim[(cikarim.tablo == "bilanco") & (cikarim.yil == 2022)]
    assert {"nominal", "sag2023"} <= set(bs.seviye)


def test_celiskili_deger_hata_veriyor(tmp_path):
    """Aynı anahtarda iki farklı değer varsa sessizce birini seçmemeli."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEST_BS"
    ws.cell(row=1, column=1, value="KALEM")
    ws.cell(row=1, column=2, value="31.12.2022 (Nominal)")
    ws.cell(row=1, column=3, value="31.12.2022 (Nominal)")
    ws.cell(row=2, column=1, value="TOPLAM VARLIKLAR")
    ws.cell(row=2, column=2, value=1_000_000.0)
    ws.cell(row=2, column=3, value=2_000_000.0)  # çelişki
    yol = tmp_path / "celiskili.xlsx"
    wb.save(yol)

    with pytest.raises(ValueError, match="çelişen"):
        calisma_kitabini_cikar(yol)
